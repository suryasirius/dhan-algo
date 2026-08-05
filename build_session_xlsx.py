import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Raw source data: (session_label, weekday -> (trades, win_rate_pct, net))
# Transcribed exactly from the user's pasted analyze_strike_momentum_sessions.py
# output (771-trade, all-weekday, all-day Strike Momentum backtest).
SESSIONS = [
    "09:15-09:45", "09:45-10:15", "10:15-10:45", "10:45-11:15",
    "11:15-11:45", "11:45-12:15", "12:15-12:45", "12:45-13:15",
    "13:15-13:45", "13:45-14:15", "14:15-14:45", "14:45-15:15", "15:15-15:30",
]

DATA = {
    "Monday": [
        (21, 47.6, 41.07), (14, 50.0, 64.14), (11, 36.4, -47.19), (8, 62.5, 22.38),
        (9, 22.2, -35.46), (7, 71.4, 22.74), (7, 42.9, 27.18), (15, 53.3, -4.74),
        (7, 57.1, -0.09), (13, 30.8, -0.54), (4, 75.0, 1.65), (7, 42.9, 22.53), (3, 33.3, -4.17),
    ],
    "Tuesday": [
        (29, 37.9, 33.12), (15, 46.7, 7.56), (8, 50.0, 12.93), (7, 28.6, -1.71),
        (12, 25.0, -53.16), (10, 40.0, -32.22), (8, 25.0, -26.22), (7, 14.3, -23.25),
        (3, 66.7, 3.06), (2, 50.0, 10.86), (4, 25.0, -4.08), (3, 100.0, 11.85), (2, 50.0, 1.59),
    ],
    "Wednesday": [
        (43, 51.2, 189.00), (30, 36.7, -92.70), (16, 56.2, 122.55), (16, 50.0, 60.72),
        (21, 42.9, -2.58), (18, 38.9, 2.16), (19, 42.1, -56.22), (17, 41.2, 0.84),
        (12, 58.3, 34.32), (15, 40.0, -12.30), (12, 25.0, -26.22), (19, 31.6, -38.73), (6, 33.3, -16.26),
    ],
    "Thursday": [
        (33, 51.5, 82.83), (8, 50.0, -17.01), (17, 35.3, -16.98), (17, 41.2, -30.48),
        (13, 23.1, 10.65), (10, 50.0, 6.60), (15, 33.3, -66.03), (11, 45.5, 13.02),
        (9, 44.4, 0.84), (9, 22.2, 3.03), (13, 46.2, -16.20), (11, 18.2, -60.93), (3, 33.3, 0.42),
    ],
    "Friday": [
        (26, 38.5, 13.32), (10, 30.0, 31.26), (10, 20.0, -42.24), (8, 37.5, -9.15),
        (10, 40.0, -23.31), (5, 80.0, 37.35), (9, 66.7, 11.73), (7, 28.6, -17.55),
        (6, 33.3, 3.45), (9, 44.4, 35.94), (10, 50.0, -3.36), (9, 66.7, 46.29), (3, 66.7, 19.71),
    ],
}

WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

# Reconstruct win counts from trades x win_rate% (rounded -- source only gave us
# the percentage, not the raw win count, so this is a faithful reconstruction).
def win_count(trades, win_pct):
    return round(trades * win_pct / 100)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "30-Min Session Analysis"

FONT_NAME = "Arial"
blue = Font(name=FONT_NAME, color="0000FF")
black = Font(name=FONT_NAME, color="000000")
bold = Font(name=FONT_NAME, bold=True)
header_fill = PatternFill("solid", fgColor="DDEBF7")
total_fill = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws["A1"] = "Strike Momentum Concept -- Win Rate / Net Points by Weekday x 30-Min Session"
ws["A1"].font = Font(name=FONT_NAME, bold=True, size=14)
ws["A2"] = ("Source: standalone 6-month backtest (test_strike_momentum.py), all weekdays, "
            "all-day entries, 771 total trades, Jan 27 - Jul 27 2026. "
            "Trades and Net Points are the raw backtest output (blue); Win Count is "
            "reconstructed from the reported win-rate percentage. Win Rate %, Avg pts, "
            "and the All Days columns are computed with formulas.")
ws["A2"].font = Font(name=FONT_NAME, italic=True, size=9, color="666666")
ws.merge_cells("A2:V2")
ws.row_dimensions[2].height = 28
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")

# Column layout: A=Session, then 4 cols per weekday (Trades, Wins, Win%, Net, Avg) -> use 5 cols/group
GROUP_COLS = ["Trades", "Wins", "Win %", "Net Pts", "Avg Pts"]
groups = WEEKDAYS + ["All Days"]

start_col = 2  # column B
group_start_col = {}
col = start_col
for g in groups:
    group_start_col[g] = col
    col += len(GROUP_COLS)

HEADER_ROW = 4
SUBHEADER_ROW = 5
FIRST_DATA_ROW = 6
LAST_DATA_ROW = FIRST_DATA_ROW + len(SESSIONS) - 1
TOTAL_ROW = LAST_DATA_ROW + 1

ws.cell(row=HEADER_ROW, column=1, value="Session").font = bold
ws.cell(row=HEADER_ROW, column=1).fill = header_fill
ws.cell(row=SUBHEADER_ROW, column=1, value="").fill = header_fill

for g in groups:
    c0 = group_start_col[g]
    ws.merge_cells(start_row=HEADER_ROW, start_column=c0, end_row=HEADER_ROW, end_column=c0 + len(GROUP_COLS) - 1)
    cell = ws.cell(row=HEADER_ROW, column=c0, value=g)
    cell.font = bold
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")
    for j, colname in enumerate(GROUP_COLS):
        sc = ws.cell(row=SUBHEADER_ROW, column=c0 + j, value=colname)
        sc.font = Font(name=FONT_NAME, bold=True, size=9)
        sc.fill = header_fill
        sc.alignment = Alignment(horizontal="center")

for i, session in enumerate(SESSIONS):
    r = FIRST_DATA_ROW + i
    ws.cell(row=r, column=1, value=session).font = black

    for wd in WEEKDAYS:
        trades, win_pct, net = DATA[wd][i]
        wins = win_count(trades, win_pct)
        c0 = group_start_col[wd]
        t_cell = ws.cell(row=r, column=c0, value=trades); t_cell.font = blue
        w_cell = ws.cell(row=r, column=c0 + 1, value=wins); w_cell.font = blue
        n_cell = ws.cell(row=r, column=c0 + 3, value=net); n_cell.font = blue; n_cell.number_format = "0.00;(0.00);-"
        wr_cell = ws.cell(row=r, column=c0 + 2,
                           value=f"=IF({get_column_letter(c0)}{r}=0,0,{get_column_letter(c0+1)}{r}/{get_column_letter(c0)}{r})")
        wr_cell.font = black
        wr_cell.number_format = "0.0%"
        avg_cell = ws.cell(row=r, column=c0 + 4,
                            value=f"=IF({get_column_letter(c0)}{r}=0,0,{get_column_letter(c0+3)}{r}/{get_column_letter(c0)}{r})")
        avg_cell.font = black
        avg_cell.number_format = "0.00;(0.00);-"

    # All Days columns -- pure formulas summing the 5 weekday columns
    c0 = group_start_col["All Days"]
    trade_cols = [get_column_letter(group_start_col[wd]) for wd in WEEKDAYS]
    win_cols = [get_column_letter(group_start_col[wd] + 1) for wd in WEEKDAYS]
    net_cols = [get_column_letter(group_start_col[wd] + 3) for wd in WEEKDAYS]

    t_formula = "=" + "+".join(f"{c}{r}" for c in trade_cols)
    w_formula = "=" + "+".join(f"{c}{r}" for c in win_cols)
    n_formula = "=" + "+".join(f"{c}{r}" for c in net_cols)

    ws.cell(row=r, column=c0, value=t_formula).font = black
    ws.cell(row=r, column=c0 + 1, value=w_formula).font = black
    ws.cell(row=r, column=c0 + 3, value=n_formula).font = black
    ws.cell(row=r, column=c0 + 3).number_format = "0.00;(0.00);-"
    ws.cell(row=r, column=c0 + 2,
            value=f"=IF({get_column_letter(c0)}{r}=0,0,{get_column_letter(c0+1)}{r}/{get_column_letter(c0)}{r})").number_format = "0.0%"
    ws.cell(row=r, column=c0 + 4,
            value=f"=IF({get_column_letter(c0)}{r}=0,0,{get_column_letter(c0+3)}{r}/{get_column_letter(c0)}{r})").number_format = "0.00;(0.00);-"
    for j in range(len(GROUP_COLS)):
        ws.cell(row=r, column=c0 + j).font = black

    for c in range(1, col):
        ws.cell(row=r, column=c).border = border

# Totals row
ws.cell(row=TOTAL_ROW, column=1, value="TOTAL").font = bold
for g in groups:
    c0 = group_start_col[g]
    t_col = get_column_letter(c0)
    w_col = get_column_letter(c0 + 1)
    n_col = get_column_letter(c0 + 3)
    ws.cell(row=TOTAL_ROW, column=c0, value=f"=SUM({t_col}{FIRST_DATA_ROW}:{t_col}{LAST_DATA_ROW})").font = bold
    ws.cell(row=TOTAL_ROW, column=c0 + 1, value=f"=SUM({w_col}{FIRST_DATA_ROW}:{w_col}{LAST_DATA_ROW})").font = bold
    net_total_cell = ws.cell(row=TOTAL_ROW, column=c0 + 3, value=f"=SUM({n_col}{FIRST_DATA_ROW}:{n_col}{LAST_DATA_ROW})")
    net_total_cell.font = bold
    net_total_cell.number_format = "0.00;(0.00);-"
    wr_total = ws.cell(row=TOTAL_ROW, column=c0 + 2,
                        value=f"=IF({t_col}{TOTAL_ROW}=0,0,{w_col}{TOTAL_ROW}/{t_col}{TOTAL_ROW})")
    wr_total.font = bold
    wr_total.number_format = "0.0%"
    avg_total = ws.cell(row=TOTAL_ROW, column=c0 + 4,
                         value=f"=IF({t_col}{TOTAL_ROW}=0,0,{n_col}{TOTAL_ROW}/{t_col}{TOTAL_ROW})")
    avg_total.font = bold
    avg_total.number_format = "0.00;(0.00);-"
    for j in range(len(GROUP_COLS)):
        ws.cell(row=TOTAL_ROW, column=c0 + j).fill = total_fill
        ws.cell(row=TOTAL_ROW, column=c0 + j).border = border
ws.cell(row=TOTAL_ROW, column=1).fill = total_fill
ws.cell(row=TOTAL_ROW, column=1).border = border

# Column widths
ws.column_dimensions["A"].width = 14
for c in range(2, col):
    ws.column_dimensions[get_column_letter(c)].width = 9

ws.freeze_panes = "B6"

wb.save("Strike_Momentum_Session_Analysis.xlsx")
print("Saved.")
